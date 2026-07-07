"""Geração de contratos genéricos em DOCX e PDF.

DOCX é gerado via python-docx.
PDF é gerado diretamente via reportlab (sem conversão DOCX→PDF,
sem dependência de LibreOffice ou Microsoft Word).
"""

from __future__ import annotations

import hashlib
import io
from datetime import date
from typing import TYPE_CHECKING

from django.core.files.base import ContentFile
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.pdfgen import canvas as rl_canvas

if TYPE_CHECKING:
    from django.contrib.auth.models import User
    from gestao_contratos.models import ContratoGerado
    from gestao_lab.models import Paciente


_CLINICA = "Associação Brasileira de Odontologia — Seção Goiás"

# Posições absolutas do bloco de assinatura no PDF (em pontos, origem no
# rodapé da página). São fixas para que o merge da assinatura eletrônica
# (services/assinatura_pdf.py) saiba exatamente onde carimbar a imagem.
PDF_MARGEM_ESQUERDA = 3.0 * cm
PDF_ASSINATURA_LARGURA = 8.5 * cm
PDF_ASSINATURA_PACIENTE_Y = 5.2 * cm
PDF_ASSINATURA_PROFISSIONAL_Y = 3.0 * cm


def gerar_e_salvar_contrato(
    paciente: "Paciente",
    tipo: str,
    observacoes_clinicas: str = "",
    profissional_nome: str = "",
    profissional_cro: str = "",
    local_assinatura: str = "Goiânia - GO",
    gerado_por: "User | None" = None,
) -> "ContratoGerado":
    """Gera DOCX + PDF, persiste ambos e retorna o ContratoGerado."""
    from gestao_contratos.models import ContratoGerado

    numero = tipo.replace("modelo_", "")
    conteudo_docx = gerar_contrato(
        paciente=paciente,
        tipo=tipo,
        profissional_nome=profissional_nome,
        profissional_cro=profissional_cro,
        local_assinatura=local_assinatura,
    )
    conteudo_pdf = gerar_pdf(
        paciente=paciente,
        tipo=tipo,
        profissional_nome=profissional_nome,
        profissional_cro=profissional_cro,
        local_assinatura=local_assinatura,
    )

    primeiro = paciente.nome.split()[0].lower()
    base = f"contrato_modelo_{numero}_{primeiro}_{paciente.id_dental}"

    versao = ContratoGerado.objects.filter(paciente=paciente, tipo=tipo).count() + 1

    contrato = ContratoGerado(
        paciente=paciente,
        tipo=tipo,
        observacoes_clinicas=observacoes_clinicas,
        profissional_nome=profissional_nome,
        profissional_cro=profissional_cro,
        local_assinatura=local_assinatura,
        gerado_por=gerado_por,
        status="gerado",
        versao=versao,
        hash_sha256=hashlib.sha256(conteudo_pdf).hexdigest(),
    )
    contrato.arquivo.save(f"{base}.docx", ContentFile(conteudo_docx), save=False)
    contrato.arquivo_pdf.save(f"{base}.pdf", ContentFile(conteudo_pdf), save=False)
    contrato.save()
    return contrato


def gerar_contrato(
    paciente: "Paciente",
    tipo: str,
    observacoes_clinicas: str = "",
    profissional_nome: str = "",
    profissional_cro: str = "",
    local_assinatura: str = "Goiânia - GO",
) -> bytes:
    """Gera e retorna os bytes DOCX com dados pessoais do paciente."""
    numero = tipo.replace("modelo_", "")
    return _gerar_docx(
        numero_modelo=numero,
        paciente=paciente,
        profissional_nome=profissional_nome,
        profissional_cro=profissional_cro,
        local_assinatura=local_assinatura,
    )


def gerar_pdf(
    paciente: "Paciente",
    tipo: str,
    profissional_nome: str = "",
    profissional_cro: str = "",
    local_assinatura: str = "Goiânia - GO",
) -> bytes:
    """Gera e retorna os bytes PDF com dados pessoais do paciente."""
    numero = tipo.replace("modelo_", "")
    return _gerar_pdf(
        numero_modelo=numero,
        paciente=paciente,
        profissional_nome=profissional_nome,
        profissional_cro=profissional_cro,
        local_assinatura=local_assinatura,
    )


def obter_melhor_pdf_bytes(contrato: "ContratoGerado") -> bytes | None:
    """Retorna os bytes do melhor PDF já salvo do contrato.

    Prioriza o PDF assinado (arquivo_pdf_assinado) sobre o original
    (arquivo_pdf) — usado por todos os canais de envio (Dental Office,
    e-mail, WhatsApp) para garantir que, uma vez assinado, é sempre a
    versão assinada que circula. Não regenera na hora: retorna None se
    nenhum dos dois arquivos existir ou puder ser lido, cabendo ao
    chamador decidir se regenera via gerar_pdf() (sempre sem assinatura).
    """
    import logging

    logger = logging.getLogger(__name__)

    for campo, rotulo in (
        (contrato.arquivo_pdf_assinado, "assinado"),
        (contrato.arquivo_pdf, "original"),
    ):
        if not campo:
            continue
        try:
            campo.open("rb")
            conteudo = campo.read()
            campo.close()
            return conteudo
        except Exception as exc:
            logger.warning(
                "obter_melhor_pdf_bytes: falha ao ler PDF %s (contrato_pk=%s): %s",
                rotulo,
                contrato.pk,
                exc,
            )
    return None


# ── Geração DOCX ─────────────────────────────────────────────────────────────


def _gerar_docx(
    numero_modelo: str,
    paciente: "Paciente",
    profissional_nome: str,
    profissional_cro: str,
    local_assinatura: str,
) -> bytes:
    doc = Document()

    for secao in doc.sections:
        secao.top_margin = Cm(2.5)
        secao.bottom_margin = Cm(2.5)
        secao.left_margin = Cm(3.0)
        secao.right_margin = Cm(2.5)

    # ── Cabeçalho ──────────────────────────────────────────────────────
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(_CLINICA.upper())
    r.bold = True
    r.font.size = Pt(10)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"MODELO {numero_modelo} — DECLARAÇÃO DE CONSENTIMENTO")
    r.bold = True
    r.font.size = Pt(13)

    doc.add_paragraph()

    # ── Dados pessoais ─────────────────────────────────────────────────
    _titulo_secao(doc, "DADOS DO PACIENTE")
    _campo(doc, "Nome completo", paciente.nome)
    _campo(doc, "Data de nascimento", _fmt_data(paciente.data_nascimento))

    if paciente.cpf:
        _campo(doc, "CPF", paciente.cpf)
    elif paciente.rg:
        _campo(doc, "RG", paciente.rg)

    endereco = _fmt_endereco(paciente)
    if endereco:
        _campo(doc, "Endereço", endereco)

    # ── Responsável legal ──────────────────────────────────────────────
    if paciente.nome_responsavel:
        doc.add_paragraph()
        _titulo_secao(doc, "RESPONSÁVEL LEGAL")
        _campo(doc, "Nome", paciente.nome_responsavel)
        if paciente.cpf_responsavel:
            _campo(doc, "CPF", paciente.cpf_responsavel)

    doc.add_paragraph()

    # ── Declaração ─────────────────────────────────────────────────────
    p = doc.add_paragraph(
        "Declaro que fui devidamente informado(a) e autorizo a realização "
        "do procedimento, estando ciente dos riscos e benefícios inerentes."
    )
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    doc.add_paragraph()

    _campo(
        doc,
        "Local e data",
        f"{local_assinatura}, {date.today().strftime('%d/%m/%Y')}",
    )

    doc.add_paragraph()
    doc.add_paragraph()

    # ── Assinaturas ────────────────────────────────────────────────────
    _linha_assinatura(doc, "Paciente ou responsável legal")

    doc.add_paragraph()

    prof = profissional_nome or "Profissional responsável"
    if profissional_cro:
        prof += f" — CRO: {profissional_cro}"
    _linha_assinatura(doc, prof)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _titulo_secao(doc: Document, texto: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(texto)
    r.bold = True
    r.font.size = Pt(10)
    p.paragraph_format.space_after = Pt(4)


def _campo(doc: Document, rotulo: str, valor: str) -> None:
    p = doc.add_paragraph()
    r_rotulo = p.add_run(f"{rotulo}: ")
    r_rotulo.bold = True
    r_rotulo.font.size = Pt(11)
    r_valor = p.add_run(valor or "—")
    r_valor.font.size = Pt(11)
    p.paragraph_format.space_after = Pt(2)


def _linha_assinatura(doc: Document, descricao: str) -> None:
    p = doc.add_paragraph("_" * 52)
    p.paragraph_format.space_after = Pt(2)
    p2 = doc.add_paragraph(descricao)
    p2.paragraph_format.space_after = Pt(6)


# ── Geração PDF (reportlab — sem dependência externa) ─────────────────────────


def _gerar_pdf(
    numero_modelo: str,
    paciente: "Paciente",
    profissional_nome: str,
    profissional_cro: str,
    local_assinatura: str,
) -> bytes:
    buf = io.BytesIO()
    larg, alt = A4
    c = rl_canvas.Canvas(buf, pagesize=A4)

    ml = PDF_MARGEM_ESQUERDA
    mr = larg - 2.5 * cm
    uw = mr - ml
    y: list[float] = [alt - 2.5 * cm]  # lista para mutabilidade em closures

    def nl(pts: float = 16.0) -> None:
        y[0] -= pts

    def _linhas(texto: str, fonte: str, sz: int, max_larg: float) -> list[str]:
        palavras = texto.split()
        resultado: list[str] = []
        linha = ""
        for p in palavras:
            cand = f"{linha} {p}".strip()
            if c.stringWidth(cand, fonte, sz) <= max_larg:
                linha = cand
            else:
                if linha:
                    resultado.append(linha)
                linha = p
        if linha:
            resultado.append(linha)
        return resultado or [""]

    def centro(texto: str, sz: int, bold: bool = False) -> None:
        c.setFont("Helvetica-Bold" if bold else "Helvetica", sz)
        c.drawCentredString(larg / 2, y[0], texto)
        nl(sz + 5)

    def esq(texto: str, sz: int, bold: bool = False) -> None:
        c.setFont("Helvetica-Bold" if bold else "Helvetica", sz)
        c.drawString(ml, y[0], texto)
        nl(sz + 4)

    def campo(rotulo: str, valor: str) -> None:
        sz = 11
        prefixo = f"{rotulo}: "
        c.setFont("Helvetica-Bold", sz)
        pw = c.stringWidth(prefixo, "Helvetica-Bold", sz)
        c.drawString(ml, y[0], prefixo)
        c.setFont("Helvetica", sz)
        texto = valor or "—"
        if c.stringWidth(texto, "Helvetica", sz) <= uw - pw:
            c.drawString(ml + pw, y[0], texto)
            nl(sz + 5)
        else:
            nl(sz + 4)
            for linha in _linhas(texto, "Helvetica", sz, uw):
                c.drawString(ml, y[0], linha)
                nl(sz + 4)
            nl(2)

    def paragrafo(texto: str, sz: int = 11) -> None:
        c.setFont("Helvetica", sz)
        for linha in _linhas(texto, "Helvetica", sz, uw):
            c.drawString(ml, y[0], linha)
            nl(sz + 4)

    def bloco_assinatura(y_linha: float, descricao: str) -> None:
        """Linha de assinatura em posição fixa, ancorada ao rodapé."""
        c.setLineWidth(0.7)
        c.line(ml, y_linha, ml + PDF_ASSINATURA_LARGURA, y_linha)
        c.setFont("Helvetica", 10)
        c.drawString(ml, y_linha - 13, descricao)

    # ── Cabeçalho ─────────────────────────────────────────────────────
    centro(_CLINICA.upper(), 10, bold=True)
    nl(4)
    centro(
        f"MODELO {numero_modelo} — DECLARAÇÃO DE CONSENTIMENTO",
        13,
        bold=True,
    )
    nl(10)

    # ── Dados pessoais ─────────────────────────────────────────────────
    esq("DADOS DO PACIENTE", 10, bold=True)
    nl(2)
    campo("Nome completo", paciente.nome)
    campo("Data de nascimento", _fmt_data(paciente.data_nascimento))
    if paciente.cpf:
        campo("CPF", paciente.cpf)
    elif paciente.rg:
        campo("RG", paciente.rg)
    endereco = _fmt_endereco(paciente)
    if endereco:
        campo("Endereço", endereco)

    # ── Responsável legal ──────────────────────────────────────────────
    if paciente.nome_responsavel:
        nl(6)
        esq("RESPONSÁVEL LEGAL", 10, bold=True)
        nl(2)
        campo("Nome", paciente.nome_responsavel)
        if paciente.cpf_responsavel:
            campo("CPF", paciente.cpf_responsavel)

    nl(10)

    # ── Declaração ─────────────────────────────────────────────────────
    paragrafo(
        "Declaro que fui devidamente informado(a) e autorizo a realização "
        "do procedimento, estando ciente dos riscos e benefícios inerentes."
    )
    nl(6)
    campo(
        "Local e data",
        f"{local_assinatura}, {date.today().strftime('%d/%m/%Y')}",
    )
    # ── Assinaturas (posição fixa, ancoradas ao rodapé) ────────────────
    bloco_assinatura(PDF_ASSINATURA_PACIENTE_Y, "Paciente ou responsável legal")
    prof = profissional_nome or "Profissional responsável"
    if profissional_cro:
        prof += f" — CRO: {profissional_cro}"
    bloco_assinatura(PDF_ASSINATURA_PROFISSIONAL_Y, prof)

    c.save()
    buf.seek(0)
    return buf.read()


# ── Helpers compartilhados ────────────────────────────────────────────────────


def _fmt_data(d: date | None) -> str:
    return d.strftime("%d/%m/%Y") if d else "—"


def _fmt_endereco(paciente: "Paciente") -> str:
    partes: list[str] = []
    logradouro = paciente.endereco_logradouro or ""
    if paciente.endereco_numero:
        logradouro = f"{logradouro}, {paciente.endereco_numero}"
    if paciente.endereco_complemento:
        logradouro = f"{logradouro} — {paciente.endereco_complemento}"
    if logradouro:
        partes.append(logradouro)
    if paciente.endereco_bairro:
        partes.append(paciente.endereco_bairro)
    cidade_estado = " - ".join(
        p for p in [paciente.endereco_cidade, paciente.endereco_estado] if p
    )
    if cidade_estado:
        partes.append(cidade_estado)
    if paciente.endereco_cep:
        partes.append(f"CEP {paciente.endereco_cep}")
    return " | ".join(partes)
